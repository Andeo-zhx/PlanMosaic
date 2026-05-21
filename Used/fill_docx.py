from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import copy

doc = Document(r'd:\Trae CN\Projects\PlanMosaic\Crouse Work\开题文档模板.docx')

# Helper: clear a paragraph and set new text
def set_paragraph_text(p, text):
    """Set paragraph text, preserving first run's formatting"""
    for run in p.runs:
        run.text = ''
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)

def add_content_para(doc, ref_p, text, bold_prefix=None):
    """Add content to a paragraph, optionally with bold prefix"""
    for run in ref_p.runs:
        run.text = ''
    if bold_prefix and ref_p.runs:
        ref_p.runs[0].text = bold_prefix
        ref_p.runs[0].bold = True
        remaining = text
        if len(ref_p.runs) > 1:
            ref_p.runs[1].text = remaining
        else:
            ref_p.add_run(remaining)
    elif ref_p.runs:
        ref_p.runs[0].text = text
    return ref_p

# ===== 1. 学号 + 姓名 (paragraph [2]) =====
p2 = doc.paragraphs[2]
for run in p2.runs:
    run.text = ''
p2.runs[0].text = '学号: （组长）322010XXXX  姓名: 朱宏轩    学号: 322010XXXX  姓名: 林盛元    学号: 322010XXXX  姓名: 孟想'

# ===== 2. 是否参加展示 (paragraph [3]) =====
p3 = doc.paragraphs[3]
p3.clear()
p3.add_run('是否参加大作业展示（是/否）：是   ✓')

# ===== 3. 技术现状 (paragraphs [7][8][9]) =====
tech_status = """当前，AI Agent（智能体）领域主要存在两种主流技术范式：基于提示工程的手工ReAct范式与基于模型原生能力的Function Calling范式。二者在技术路线、对模型的要求以及适用场景上存在显著差异。

（一）ReAct范式（Reasoning + Acting）。由Yao等人于2022年在ICLR 2023发表的论文《ReAct: Synergizing Reasoning and Acting in Language Models》中首次提出。其核心思想是让LLM在解决问题的过程中交替生成"思考轨迹"（Thought）和"行动指令"（Action），并通过外部环境反馈（Observation）形成推理闭环。ReAct的本质优势在于：① 思维的透明性——每一步推理过程都显式可读，便于调试和验证；② 对模型底座要求低——不需要模型本身支持function calling格式输出，任何文本生成模型均可使用；③ 可通过少量示例（Few-shot）激活协同能力，在HotpotQA、Fever等知识密集任务中较纯推理基线减少了事实幻觉和错误传播。

（二）Function Calling范式。以OpenAI于2023年6月推出的Function Calling API为代表，随后DeepSeek、Qwen、Claude等主流模型服务商均跟进实现。该范式通过预定义JSON Schema格式的工具描述列表，由模型在推理过程中自主判断是否需要调用工具，并直接输出结构化的函数名和参数JSON。相比ReAct，Function Calling具有解析稳定性高（不需要正则匹配自由文本）、调用延迟低（单次API调用可完成判断+参数构造）、多工具并行调用等优势。根据AgentArch基准测试（2025），在生产环境中Function Calling Agent在响应速度上通常优于文本型ReAct Agent。

（三）混合方案与工业实践。当前业界主流Agent框架（如LangChain、AutoGPT、CrewAI等）均在两种范式之上构建了高层抽象。然而，选题明确要求不允许使用这些高层框架——这一约束恰好促使我们从原理层面深入理解Agent的工作机制。值得注意的是，Function Calling在底层仍然遵循"推理→选择工具→构造参数→执行→接收反馈"的ReAct逻辑，只是将Thought和Action的生成内化到了模型的参数化推理中，而非显式的文本输出。因此，将隐式的Function Calling流程回放为显式的ReAct文本轨迹，不仅能够满足课程对ReAct范式的要求，还能够在不牺牲交互体验的前提下深入理解两种范式的等价性。

（四）工具生态。Agent的能力边界本质上由工具决定。当前主流工具类型包括：知识检索（维基百科、搜索引擎API）、数学计算（Python eval/NumPy）、文件操作（本地读写）、代码执行（沙箱环境）、网络爬取等。近年来，研究者也在探索将小模型应用于工具本身的智能化，如利用回归模型预估任务耗时、利用文本分类模型进行意图路由等，使工具本身也具备一定的"智能"属性。"""

p7 = doc.paragraphs[7]
p7.clear()
p7.add_run(tech_status)

# ===== 4. 核心原理研究 (paragraphs [11][12][13]) =====
core_principles = """本项目的核心逻辑可以分为四个层次来理解：

（一）Prompt工程——约束LLM的行为边界。System Prompt是整个Agent的"操作系统"。通过精心设计的提示词，我们为LLM定义了身份角色（日程规划助手/战略顾问）、行为规范（温暖亲切的语气、需要确认的操作类型）、能力边界（可用工具列表及调用规范）以及输出格式要求。对于ReAct范式，Prompt中需要通过Few-shot示例教导模型按"Thought→Action→Observation"格式交替输出。而在Function Calling范式中，Prompt更多承担"何时应该调用工具"以及"如何理解工具返回结果"的语义指导角色。本项目在既有Function Calling System Prompt基础上，通过解析对话历史中的tool_calls序列将其逆向翻译为ReAct格式文本，这一转换过程本身要求对两种Prompt范式都有透彻理解。

（二）工具定义与注册——Agent的动作空间。每个工具由两部分组成：① JSON Schema格式的接口描述（函数名、参数类型、参数含义、必填字段），供LLM理解工具的语义和适用范围；② 执行函数，负责实际调用外部API或执行本地计算。工具描述的质量直接影响LLM的工具选择准确率——描述需要足够清晰具体但又不冗余。本项目现有的14+工具（日程CRUD、冲突检测、深度分析等）已建立了成熟的Schema定义模式，新增工具只需遵循相同规范即可无缝接入。

（三）解析与执行循环——Agent的核心引擎。这是整个框架最关键的部分。在Function Calling模式下，LLM返回的响应中可能包含一个tool_calls数组，每个元素包含工具名和参数JSON。后端解析该数组后逐项执行，将每个工具的返回结果以tool角色的消息形式追加到对话上下文中，然后再次调用LLM——此时LLM可以看到工具的执行结果并决定下一步操作。循环终止的条件包括：① LLM不再返回tool_calls（认为任务已完成）；② 达到最大递归深度限制（防止死循环）；③ LLM返回了需要用户确认的Proposal。在ReAct模式下，此循环需要额外的一层文本解析器——从LLM的原始文本输出中通过正则或状态机提取"Action: tool_name(args)"片段。无论是哪种模式，执行循环的质量决定了Agent的健壮性，需要充分考虑JSON解析失败的回退策略、工具执行异常时的错误信息反馈、以及上下文窗口的Token控制。

（四）对话历史与上下文管理。Agent的每一次交互都在消耗有限上下文窗口（DeepSeek约64K、Qwen约32K tokens）。有效的上下文管理策略包括：① 按角色保留最近N轮完整对话；② 对过长的工具返回结果进行摘要压缩；③ 利用tiktoken等库精确计算token数以做出裁剪决策。本项目在多端架构中通过DataStore/JSON文件实现了对话历史的本地持久化，Python端可在此基础上增强Token计数与智能截断能力。"""

p11 = doc.paragraphs[11]
p11.clear()
p11.add_run(core_principles)

# ===== 5. 团队分工表 (Table 0) =====
table0 = doc.tables[0]
# Row 1: 朱宏轩
table0.rows[1].cells[0].text = '朱宏轩（组长）'
table0.rows[1].cells[1].text = '系统架构 + Agent核心逻辑'
table0.rows[1].cells[2].text = '① ReAct转录功能设计与实现（对话历史→ReAct文本的转换引擎）；② Python后端服务架构（FastAPI搭建、与三端前端的API对接）；③ System Prompt工程（ReAct格式约束与现有助手风格的融合）；④ 整体项目集成与质量把控'

# Row 2: 林盛元
table0.rows[2].cells[0].text = '林盛元'
table0.rows[2].cells[1].text = '工具开发 + 数据与模型'
table0.rows[2].cells[2].text = '① 新增工具实现：网络搜索评估工具（web_search_evaluate——对接Bing/SerpAPI）、任务时间估算工具（estimate_task_time——规则版+统计版）；② 可选Python微服务：时间估算小模型训练（scikit-learn RandomForest）；③ 日程健康度检查功能（analyze工具扩展）；④ 多步调用演示案例制作'

# Row 3: 孟想
table0.rows[3].cells[0].text = '孟想'
table0.rows[3].cells[1].text = '前端对接 + 文档与演示'
table0.rows[3].cells[2].text = '① 三端前端（Electron/Android/Uni-app）对接Python后端（API地址配置、接口联调）；② ReAct转录按钮的UI实现（桌面端弹窗展示+复制、Android端菜单选项、Uni-app端按钮）；③ 开题文档撰写与技术调研报告；④ 演示视频录制与答辩PPT制作'

# ===== 6. 任务进度表 (Table 1) =====
table1 = doc.tables[1]
table1.rows[1].cells[1].text = '【需求确认与开题文档】\n· 三人共同讨论确定技术方案与改进方向\n· 朱宏轩：完成开题文档技术背景调研部分，搭建Python Flask/FastAPI项目骨架\n· 林盛元：调研网络搜索API选型（Bing/SERP/DuckDuckGo），撰写工具设计文档\n· 孟想：调研现有三端前端代码中API调用位置，撰写前端改动方案\n· 交付物：开题文档.docx、项目技术方案定稿'

table1.rows[2].cells[1].text = '【核心功能开发】\n· 朱宏轩：实现ReAct转录引擎（对话历史解析+格式转换），完成Python后端HTTP API（/api/agent-chat、/api/react-log等端点），System Prompt调整\n· 林盛元：实现web_search_evaluate工具（对接搜索API），实现estimate_task_time工具（规则版+历史数据统计版），扩展analyze工具health_check维度\n· 孟想：三端前端添加ReAct转录按钮UI，前后端联调测试，开始撰写演示脚本\n· 交付物：可运行的Python Agent后端 + 新工具集成 + 前端对接完成'

table1.rows[3].cells[1].text = '【集成测试与提交】\n· 全员：端到端测试（至少3个多步调用案例验证），Bug修复\n· 朱宏轩：代码审查与最终集成，确保三端正常运行\n· 林盛元：多步调用演示案例录制（含ReAct转录展示），性能优化\n· 孟想：录制演示视频（5分钟展示），完成开题文档终稿，整理提交材料（代码+视频+文档）\n· 交付物：最终可运行项目、演示视频.mp4、开题文档终稿.docx'

# ===== 7. 补充更多技术细节到 empty paragraphs =====
# Paragraph [8] - 补充技术现状
p8 = doc.paragraphs[8]
p8.clear()
additional_tech = """（五）任务时间估算的智能化趋势。传统Agent工具多为"被动响应式"——即严格按预设逻辑执行，不具备自适应能力。近年来，研究者开始探索将轻量级机器学习模型嵌入工具层，使工具本身具备一定的"智能"属性。例如，在任务管理场景中，用户经常面临"不知道某项任务需要多长时间"的困境——这直接影响了日程排期的合理性。通过收集用户历史任务数据（任务名称、类别、预计耗时、实际耗时），可以训练一个回归模型（如RandomForest或XGBoost）来预估新任务的合理耗时。此类"智能工具"的实现位于Agent框架的外围，不改变Agent的核心循环，但显著提升了用户体验和规划的准确性。Python生态中scikit-learn、pandas等库为快速原型开发提供了良好支持。

（六）网络搜索作为Agent的关键能力。LLM存在固有的知识截止日期限制（通常截止于训练数据时间点），且无法获取实时信息。将网络搜索能力赋予Agent，使其能够在遇到超出知识范围的问题时主动检索最新信息，是Agent从"静态知识库"升级为"动态信息处理系统"的关键一步。主流实现方案包括调用Bing Web Search API、SerpAPI（Google搜索结果）、DuckDuckGo Instant Answer API等。搜索工具返回的结果需经Agent二次加工（摘要、对比、筛选），而非原样呈现给用户——这体现了Agent作为"信息中间件"的核心价值。"""
p8.add_run(additional_tech)

# Paragraph [12] - 补充核心原理
p12 = doc.paragraphs[12]
p12.clear()
more_principles = """（五）安全确认机制的设计哲学。Agent的自主性是一把双刃剑——赋予LLM直接操作本地文件或修改用户数据的能力，必须辅以相应的安全护栏。本项目采用的Proposal机制是一种"建议-确认"模式：对于危险操作（删除日程、批量修改等），Agent不直接执行，而是生成一个结构化的Proposal对象，包含操作类型、影响范围和修改理由，通过前端UI渲染为可视化的变更预览（如左右对比卡片），等待用户显式确认后才执行。这一设计将"LLM的自主决策"与"用户的最终审批权"有机结合，是在Agent能力与系统安全之间取得平衡的关键工程实践。

（六）三端统一后端架构的优势。与从头搭建单体Python应用不同，本项目已具备Electron桌面端、Android原生端和Uni-app跨平台端三个成熟前端。新增的Python Agent后端作为一个独立的HTTP微服务运行，三端前端仅需将API请求指向该服务即可获得增强后的AI能力。这种架构具备高度解耦性：前端UI的迭代与后端Agent逻辑的演进可以独立进行。Python端的引入遵循"加法而非替换"原则——server.js中现有的日程管理逻辑保持不变，Python端专注于Agent核心循环（ReAct转录、工具执行、上下文管理）和计算密集型任务（模型训练、数据分析）。"""
p12.add_run(more_principles)

# Paragraph [13] - references
p13 = doc.paragraphs[13]
p13.clear()
refs_text = """参考文献：
[1] Yao S, Zhao J, Yu D, et al. ReAct: Synergizing Reasoning and Acting in Language Models[C]. ICLR 2023 (Notable Top 5%).
[2] OpenAI. Function Calling API Documentation, 2023-2025.
[3] DeepSeek. API Documentation - Function Calling, 2024-2025.
[4] Wei J, Wang X, Schuurmans D, et al. Chain-of-Thought Prompting Elicits Reasoning in Large Language Models[C]. NeurIPS 2022.
[5] AgentArch: A Comprehensive Benchmark to Evaluate Agent Architectures in Enterprise, arXiv 2025.
[6] LeewayHertz. ReAct Agents vs Function Calling Agents: A Comparative Analysis, 2024.
[7] LangChain Documentation: Agents & Tool Use, 2024-2025."""
p13.add_run(refs_text)

# Save
out_path = r'd:\Trae CN\Projects\PlanMosaic\Crouse Work\开题文档_完整版.docx'
doc.save(out_path)
print(f'Saved to {out_path}')
